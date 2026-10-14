# -*- coding: utf-8 -*-
"""将 main_manuscript_CN.md 转成排版完整的 Word 文档（含表格、粗体、嵌入配图）。"""
import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn

MD = "papers/main_manuscript_CN.md"
OUT = "papers/main_manuscript_CN.docx"

doc = Document()

# 全局默认中文字体
def set_font(style_or_run, name_zh="宋体", name_en="Times New Roman", size=None, bold=None, color=None):
    f = style_or_run.font
    f.name = name_en
    if name_zh:
        style_or_run._element.rPr.rFonts.set(qn("w:eastAsia"), name_zh) if hasattr(style_or_run, "_element") else None
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if color: f.color.rgb = color

# 设置默认 Normal 样式
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
normal.font.size = Pt(11)

def set_eastasia(run, zh="宋体"):
    rpr = run._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), zh)

def add_runs_with_format(p, text):
    """解析 **bold** 内联粗体。"""
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            r.bold = True
            set_eastasia(r)
        else:
            r = p.add_run(part)
            set_eastasia(r)
    return p

def add_heading(text, level):
    sizes = {1: 18, 2: 15, 3: 13}
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = (level <= 2)
    r.font.size = Pt(sizes.get(level, 13))
    r.font.name = "Times New Roman"
    set_eastasia(r, "微软雅黑" if level <= 2 else "宋体")
    if level == 1:
        r.font.color.rgb = RGBColor(0x1F, 0x3B, 0x63)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
    else:
        r.font.color.rgb = RGBColor(0x17, 0x17, 0x17)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
    return p

def is_table_line(line):
    return line.strip().startswith("|") and line.strip().endswith("|")

def is_sep_line(line):
    s = line.strip()
    return s.startswith("|") and re.fullmatch(r"\|[\s:\-|]+\|", s) is not None

def parse_table_row(line):
    s = line.strip().strip("|")
    return [c.strip() for c in s.split("|")]

def add_table(rows):
    header = [parse_table_row(rows[0])]
    # rows[1] is separator
    body = [parse_table_row(r) for r in rows[2:]]
    ncols = len(header[0])
    table = doc.add_table(rows=1 + len(body), cols=ncols)
    table.style = "Light Grid Accent 1"
    # header
    for j, cell in enumerate(table.rows[0].cells):
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(header[0][j])
        r.bold = True
        set_eastasia(r)
        r.font.size = Pt(10)
    for i, row in enumerate(body):
        for j in range(ncols):
            cell = table.rows[i + 1].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            r = p.add_run(row[j] if j < len(row) else "")
            set_eastasia(r)
            r.font.size = Pt(9.5)
    return table

lines = open(MD, encoding="utf-8").read().splitlines()

i = 0
n = len(lines)
while i < n:
    line = lines[i]
    s = line.strip()

    if not s:
        i += 1
        continue

    # 水平线
    if re.fullmatch(r"-{3,}", s):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        i += 1
        continue

    # 标题
    if s.startswith("### "):
        add_heading(s[4:], 3); i += 1; continue
    if s.startswith("## "):
        add_heading(s[3:], 2); i += 1; continue
    if s.startswith("# "):
        add_heading(s[2:], 1); i += 1; continue

    # 表格
    if is_table_line(line):
        # 收集连续表格块
        block = []
        while i < n and is_table_line(lines[i]):
            block.append(lines[i]); i += 1
        # 去掉分隔行
        add_table(block)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        continue

    # 引用块 (>) —— 作为普通段落
    if s.startswith(">"):
        txt = s.lstrip("> ").strip()
        p = doc.add_paragraph()
        add_runs_with_format(p, txt)
        p.paragraph_format.left_indent = Inches(0.3)
        i += 1; continue

    # 无序列表
    if s.startswith("- ") or re.match(r"^\d+\.\s", s):
        p = doc.add_paragraph(style="List Bullet" if s.startswith("- ") else "List Number")
        txt = re.sub(r"^[-]\s+", "", s)
        txt = re.sub(r"^\d+\.\s+", "", txt)
        add_runs_with_format(p, txt)
        p.paragraph_format.space_after = Pt(2)
        i += 1; continue

    # 普通段落（含 **bold**）
    p = doc.add_paragraph()
    add_runs_with_format(p, s)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(6)
    # 正文两端对齐
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    i += 1

# ---- 嵌入配图（附录） ----
doc.add_page_break()
add_heading("附录：配图", 2)
figs = [
    ("fig1_workflow.png", "图 1. 研究设计与分析流程"),
    ("fig2_volcano_all_prot.png", "图 2. 泛肿瘤形态—蛋白质组关联火山图"),
    ("fig3_sig_counts.png", "图 3. 各模态与队列显著关联数量"),
    ("fig4_enrichment.png", "图 4. 显著蛋白命中通路富集点图"),
    ("fig5_os_cox.png", "图 5. 形态特征对总生存期的预后表现"),
    ("fig6_frozen_repro.png", "图 6. 跨生态位分解模型的可复现性"),
]
for fname, cap in figs:
    p = doc.add_paragraph()
    r = p.add_run()
    try:
        r.add_picture(f"papers/figures/{fname}", width=Inches(6.0))
    except Exception as e:
        r.add_text(f"[图片缺失: {fname}]")
        print("  ! missing:", fname, e)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp = doc.add_paragraph()
    cr = cp.add_run(cap)
    cr.bold = True
    set_eastasia(cr, "微软雅黑")
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(12)

doc.save(OUT)
print("saved ->", OUT)